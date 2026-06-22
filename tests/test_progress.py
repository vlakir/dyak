"""Тесты прогресса `generate` (T013): JSONL-контракт + бар только в TTY."""

from __future__ import annotations

import json
import os
import pty
import subprocess
import sys
from typing import TYPE_CHECKING

import pytest
from docx import Document
from openpyxl import Workbook
from typer.testing import CliRunner

from dyak.cli import app
from dyak.progress import GenerateProgress

if TYPE_CHECKING:
    from pathlib import Path

runner = CliRunner()

_HEADERS = ['Фамилия', 'Имя', 'Отчество', 'Должность']
_ROWS = [
    ['Иванов', 'Пётр', 'Семёнович', 'директор'],
    ['Петрова', 'Анна', 'Сергеевна', 'главный бухгалтер'],
]


def _make_xlsx(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(_HEADERS)
    for row in _ROWS:
        ws.append(row)
    wb.save(path)
    return path


def _make_template(path: Path, body: str) -> Path:
    doc = Document()
    doc.add_paragraph(body)
    doc.save(path)
    return path


def _jsonl(stderr: str) -> list[dict[str, object]]:
    return [
        json.loads(line)
        for line in stderr.splitlines()
        if line.strip().startswith('{')
    ]


def _invoke(tmp_path: Path, body: str, *, flag: bool) -> object:
    table = _make_xlsx(tmp_path / 'table.xlsx')
    template = _make_template(tmp_path / 'tpl.docx', body)
    args = [
        'generate',
        '--table',
        str(table),
        '--template',
        str(template),
        '--out',
        str(tmp_path / 'out'),
    ]
    if flag:
        args.append('--progress-json')
    return runner.invoke(app, args)


def test_progress_json_emits_event_stream(tmp_path: Path) -> None:
    result = _invoke(tmp_path, 'Назначить {{ Фамилия }}.', flag=True)

    assert result.exit_code == 0
    events = _jsonl(result.stderr)
    kinds = [event['event'] for event in events]
    assert kinds == ['start', 'progress', 'progress', 'done']
    assert events[0] == {'event': 'start', 'total': 2}
    assert events[-1] == {'event': 'done', 'ok': 2, 'failed': 0}
    assert events[1]['done'] == 1
    assert events[2]['done'] == 2


def test_progress_json_keeps_stdout_clean(tmp_path: Path) -> None:
    result = _invoke(tmp_path, 'Назначить {{ Фамилия }}.', flag=True)

    assert result.exit_code == 0
    # Машиночитаемый поток — только в stderr, stdout под человеческий итог.
    assert '{"event"' not in result.stdout
    assert 'Готово' in result.stdout


def test_no_flag_keeps_stderr_clean(tmp_path: Path) -> None:
    result = _invoke(tmp_path, 'Назначить {{ Фамилия }}.', flag=False)

    assert result.exit_code == 0
    # Без флага в не-TTY (тест) бар отключён, JSONL не эмитится.
    assert '{"event"' not in result.stderr


def test_progress_json_emits_error_on_failure(tmp_path: Path) -> None:
    # Неизвестная переменная → TemplateError прерывает прогон (T004).
    result = _invoke(tmp_path, 'Привет {{ Неизвестная }}.', flag=True)

    assert result.exit_code == 1
    events = _jsonl(result.stderr)
    assert events[0]['event'] == 'start'
    assert events[-1]['event'] == 'error'
    assert 'Неизвестная' in str(events[-1]['message'])


def test_generate_progress_event_order(capsys: pytest.CaptureFixture[str]) -> None:
    with GenerateProgress(2, json_events=True) as progress:
        progress.advance('a.docx')
        progress.advance('b.docx')

    err = capsys.readouterr().err
    kinds = [event['event'] for event in _jsonl(err)]
    assert kinds == ['start', 'progress', 'progress', 'done']


def test_bar_disabled_without_tty() -> None:
    # В не-TTY (pytest) живой бар не должен рисоваться.
    progress = GenerateProgress(3, json_events=False)
    assert progress._progress.disable is True


def test_bar_renders_in_tty(tmp_path: Path) -> None:
    # Под псевдотерминалом rich-бар отрисовывается (описание «Генерация»).
    code = (
        'from dyak.progress import GenerateProgress\n'
        "with GenerateProgress(2, json_events=False) as p:\n"
        "    p.advance('a'); p.advance('b')\n"
    )
    main_fd, child_fd = pty.openpty()
    env = {**os.environ, 'PYTHONPATH': 'src', 'TERM': 'xterm', 'COLUMNS': '80'}
    subprocess.run(
        [sys.executable, '-c', code],
        stderr=child_fd,
        stdout=subprocess.DEVNULL,
        env=env,
        check=True,
        close_fds=True,
    )
    os.close(child_fd)

    chunks = bytearray()
    while True:
        try:
            data = os.read(main_fd, 4096)
        except OSError:
            break
        if not data:
            break
        chunks.extend(data)
    os.close(main_fd)

    rendered = chunks.decode('utf-8', 'replace')
    assert 'Генерация' in rendered
