"""
Снимок поведения «Дьяка» на `examples/`: текст документов + отчёты `check`.

Гейт регрессии, заведённый в T033 (перевод движка в библиотеку
`chancellery`). Смысл: движок живёт теперь в чужом репозитории и обновляется
своей версией, а убедиться, что склонение не поехало, можно только по
результату — готовым документам.

Снять снимок::

    PYTHONPATH=src uv run python scripts/snapshot_examples.py /tmp/snap

Сверить с эталоном (код возврата 1 при расхождении)::

    PYTHONPATH=src uv run python scripts/snapshot_examples.py /tmp/snap
        --baseline scripts/examples_baseline.json

Эталон снят на «Дьяке» 0.3.3 (коммит 47e02eb), ДО перевода на библиотеку.
Сверка поабзацная: расхождение печатается как «было / стало» с именем
документа и номером абзаца. Побайтно сравнивать нельзя — docx это zip с
плавающими метаданными.

Раннер подменяем (`--runner module:function`): библиотека `chancellery`
гоняет тот же снимок у себя, подставляя обвязку вместо CLI «Дьяка».
"""

from __future__ import annotations

import argparse
import importlib
import json
import subprocess
import sys
from pathlib import Path
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from collections.abc import Callable

import docx

ROOT = Path(__file__).resolve().parent.parent
EX = ROOT / 'examples'

SCENARIOS = [
    (
        'plain',
        ['generate', '--table', 'employees.xlsx', '--template', 'order_template.docx'],
    ),
    (
        'nonstandard',
        [
            'generate',
            '--table',
            'employees_nonstandard.xlsx',
            '--template',
            'order_template.docx',
        ],
    ),
    (
        'ranks',
        [
            'generate',
            '--table',
            'personnel_ranks.xlsx',
            '--template',
            'rank_order_template.docx',
        ],
    ),
    (
        'military',
        [
            'generate',
            '--table',
            'military_units.xlsx',
            '--template',
            'military_order_template.docx',
        ],
    ),
]

CHECKS = [
    (
        'check_plain',
        ['check', '--table', 'employees.xlsx', '--template', 'order_template.docx'],
    ),
    (
        'check_ranks',
        [
            'check',
            '--table',
            'personnel_ranks.xlsx',
            '--template',
            'rank_order_template.docx',
        ],
    ),
    (
        'check_military',
        [
            'check',
            '--table',
            'military_units.xlsx',
            '--template',
            'military_order_template.docx',
        ],
    ),
]


def run_via_dyak_cli(args: list[str], out: Path | None) -> dict[str, object]:
    """Раннер по умолчанию: «Дьяк» через CLI. Подменяется флагом --runner."""
    cmd = [sys.executable, '-m', 'dyak', *args]
    if out is not None:
        cmd += ['--out', str(out)]
    proc = subprocess.run(
        cmd,
        cwd=EX,
        capture_output=True,
        text=True,
        env={
            'PYTHONPATH': str(ROOT / 'src'),
            'PATH': '/usr/bin:/bin',
            'LANG': 'ru_RU.UTF-8',
        },
        check=False,
    )
    return {'code': proc.returncode, 'stdout': proc.stdout, 'stderr': proc.stderr}


def scrub(text: str, out: Path | None) -> str:
    """Убрать из вывода абсолютные пути — иначе снимки двух машин не сойдутся."""
    if out is not None:
        text = text.replace(str(out), '<OUT>')
    return text.replace(str(EX), '<EXAMPLES>').replace(str(ROOT), '<ROOT>')


def read_docx(path: Path) -> list[str]:
    document = docx.Document(str(path))
    body = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                body += [p.text for p in cell.paragraphs if p.text.strip()]
    return body


def load_runner(spec: str | None) -> Callable[..., dict[str, object]]:
    """--runner module:function → вызываемое (args, out) -> dict."""
    if spec is None:
        return run_via_dyak_cli
    module_name, _, func_name = spec.partition(':')
    return getattr(importlib.import_module(module_name), func_name)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('dest', type=Path)
    parser.add_argument(
        '--runner', default=None, help='module:function вместо «Дьяка» через CLI'
    )
    parser.add_argument(
        '--baseline',
        type=Path,
        default=None,
        help='эталонный snapshot.json — сверить и выйти с кодом',
    )
    ns = parser.parse_args()

    run = load_runner(ns.runner)
    dest = ns.dest
    dest.mkdir(parents=True, exist_ok=True)
    snapshot: dict[str, object] = {}

    for name, args in SCENARIOS:
        out = dest / 'out' / name
        result = run(args, out)
        result['stdout'] = scrub(str(result['stdout']), out)
        result['stderr'] = scrub(str(result['stderr']), out)
        docs = {}
        if out.exists():
            for f in sorted(out.glob('*.docx')):
                docs[f.name] = read_docx(f)
        snapshot[name] = {'run': result, 'documents': docs}

    for name, args in CHECKS:
        result = run(args, None)
        result['stdout'] = scrub(str(result['stdout']), None)
        result['stderr'] = scrub(str(result['stderr']), None)
        snapshot[name] = result

    (dest / 'snapshot.json').write_text(
        json.dumps(snapshot, ensure_ascii=False, indent=2), encoding='utf-8'
    )

    total = sum(
        len(v['documents'])
        for k, v in snapshot.items()
        if isinstance(v, dict) and 'documents' in v
    )
    print(f'сценариев: {len(SCENARIOS)}, документов: {total}')
    for name in (n for n, _ in SCENARIOS):
        entry = snapshot[name]
        print(
            f'  {name}: код {entry["run"]["code"]}, '
            f'{len(entry["documents"])} документов'
        )
    for name in (n for n, _ in CHECKS):
        print(f'  {name}: код {snapshot[name]["code"]}')

    if ns.baseline is not None:
        expected = json.loads(ns.baseline.read_text(encoding='utf-8'))
        diffs = compare(expected, snapshot)
        if diffs:
            print(f'\nРАСХОЖДЕНИЙ: {len(diffs)}')
            for line in diffs[:40]:
                print('  ', line)
            sys.exit(1)
        print('\nСверка с эталоном: совпадает полностью.')


def compare(expected: dict, actual: dict) -> list[str]:
    """Сравнить два снимка. Возвращает список расхождений человеческим текстом."""
    diffs: list[str] = []
    for key, exp in expected.items():
        act = actual.get(key)
        if act is None:
            diffs.append(f'{key}: сценарий отсутствует в новом снимке')
            continue
        if exp.get('code', exp.get('run', {}).get('code')) != act.get(
            'code', act.get('run', {}).get('code')
        ):
            diffs.append(f'{key}: код возврата разошёлся')
        if 'documents' in exp:
            for doc, paragraphs in exp['documents'].items():
                got = act['documents'].get(doc)
                if got is None:
                    diffs.append(f'{key}/{doc}: документ не порождён')
                    continue
                for i, (a, b) in enumerate(zip(paragraphs, got, strict=False)):
                    if a != b:
                        diffs.append(
                            f'{key}/{doc} абзац {i}:\n      было: {a}\n      стало: {b}'
                        )
                if len(paragraphs) != len(got):
                    diffs.append(
                        f'{key}/{doc}: абзацев было {len(paragraphs)}, стало {len(got)}'
                    )
            diffs.extend(
                f'{key}/{doc}: лишний документ'
                for doc in act['documents']
                if doc not in exp['documents']
            )
        elif exp.get('stdout') != act.get('stdout'):
            diffs.append(
                f'{key}: отчёт разошёлся\n'
                f'      было: {exp.get("stdout", "").strip()}\n'
                f'      стало: {act.get("stdout", "").strip()}'
            )
    return diffs


if __name__ == '__main__':
    main()
